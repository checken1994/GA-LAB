from __future__ import annotations

import argparse
import csv
import hashlib
import json
from collections import Counter
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

REQUIRED = {
    "schema_version", "question_id", "question", "candidate_answer", "legacy_source_url", "proposed_canonical_source_url",
    "source_fetch_status", "source_http_status", "source_fetch_timestamp", "source_content_sha256", "source_cache_path",
    "source_canonical_url_observed", "gold_chunk_id", "gold_chunk_text", "gold_chunk_sha256", "gold_evidence_quote",
    "quote_contained_in_fetched_source", "alignment_status", "temporal_scope", "review_status", "reviewer_id",
    "reviewer_provenance", "reviewed_at", "review_method", "review_notes", "human_review_required", "gold_promotion",
    "eligible_for_ragas", "candidate_only",
}


def digest(value: str | bytes) -> str:
    data = value if isinstance(value, bytes) else value.encode("utf-8")
    return "sha256:" + hashlib.sha256(data).hexdigest()


def load_jsonl(path: Path) -> list[dict]:
    return [json.loads(line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]


def validate(args: argparse.Namespace) -> dict:
    rows = load_jsonl(args.jsonl)
    issues: list[str] = []
    ids = [row.get("question_id") for row in rows]
    if len(rows) != 50:
        issues.append(f"JSONL_ROW_COUNT:{len(rows)}")
    if len(set(ids)) != len(ids):
        issues.append("JSONL_DUPLICATE_OR_EMPTY_ID")
    for index, row in enumerate(rows, 1):
        missing = sorted(REQUIRED - set(row))
        if missing:
            issues.append(f"ROW_{index}_MISSING:{','.join(missing)}")
        quote = str(row.get("gold_evidence_quote") or "")
        chunk = str(row.get("gold_chunk_text") or "")
        if row.get("quote_contained_in_fetched_source") and (not quote or quote not in chunk):
            issues.append(f"ROW_{index}_QUOTE_FLAG_FALSE")
        if quote and quote not in chunk:
            issues.append(f"ROW_{index}_QUOTE_NOT_CONTAINED")
        if chunk and row.get("gold_chunk_sha256") != digest(chunk):
            issues.append(f"ROW_{index}_CHUNK_HASH_INVALID")
        if row.get("eligible_for_ragas") is not False:
            issues.append(f"ROW_{index}_RAGAS_ELIGIBILITY_NOT_FAIL_CLOSED")
        if row.get("gold_promotion") != "BLOCKED_PENDING_HUMAN_REVIEW":
            issues.append(f"ROW_{index}_PROMOTION_NOT_BLOCKED")
        if row.get("human_review_required") is not True or row.get("candidate_only") is not True:
            issues.append(f"ROW_{index}_REVIEW_FLAGS_INVALID")
    fields = list(rows[0].keys()) if rows else []
    if args.csv.exists():
        with args.csv.open(encoding="utf-8-sig", newline="") as handle:
            reader = csv.DictReader(handle)
            if reader.fieldnames != fields:
                issues.append("CSV_HEADER_MISMATCH")
            csv_rows = list(reader)
        if len(csv_rows) != len(rows):
            issues.append(f"CSV_ROW_COUNT:{len(csv_rows)}")
    else:
        issues.append("CSV_MISSING")
    if args.xlsx.exists():
        workbook = load_workbook(args.xlsx, read_only=True, data_only=False)
        if "Gold Candidate v2" not in workbook.sheetnames:
            issues.append("XLSX_SHEET_MISSING")
        else:
            sheet = workbook["Gold Candidate v2"]
            values = list(sheet.iter_rows(values_only=True))
            header = [str(value) if value is not None else None for value in values[0]] if values else []
            if header != fields:
                issues.append("XLSX_HEADER_MISMATCH")
            if len(values) - 1 != len(rows):
                issues.append(f"XLSX_ROW_COUNT:{len(values)-1}")
    else:
        issues.append("XLSX_MISSING")
    result = {
        "schema_version": "phase3-gold-v2-validator",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "jsonl": str(args.jsonl),
        "csv": str(args.csv),
        "xlsx": str(args.xlsx),
        "jsonl_sha256": digest(args.jsonl.read_bytes()),
        "rows": len(rows),
        "status_counts": dict(Counter(row.get("review_status") for row in rows)),
        "quote_contained_count": sum(bool(row.get("quote_contained_in_fetched_source")) for row in rows),
        "eligible_for_ragas_count": sum(bool(row.get("eligible_for_ragas")) for row in rows),
        "issues": issues,
        "schema_validation": "PASS_WITHIN_SCOPE" if not issues else "FAIL",
        "gold_promotion_gate": "BLOCKED_PENDING_INDEPENDENT_REVIEW",
        "note": "Structural validation PASS does not promote candidate rows to verified Gold.",
    }
    args.report.parent.mkdir(parents=True, exist_ok=True)
    args.report.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--jsonl", type=Path, required=True)
    parser.add_argument("--csv", type=Path, required=True)
    parser.add_argument("--xlsx", type=Path, required=True)
    parser.add_argument("--report", type=Path, required=True)
    args = parser.parse_args()
    result = validate(args)
    print(json.dumps(result, ensure_ascii=False, indent=2))
    raise SystemExit(0 if result["schema_validation"] == "PASS_WITHIN_SCOPE" else 1)


if __name__ == "__main__":
    main()
