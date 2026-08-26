from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import time
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import quote, unquote, urlparse

import requests
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STATIC_WIKI_ROWS = {
    "CH-0005": ["spaced repetition", "flashcards"],
    "CH-0006": ["điện biên phủ", "1954"],
    "CH-0008": ["el niño", "climate"],
    "CH-0011": ["large language model", "llm"],
    "CH-0017": ["pop music", "1950"],
    "CH-0018": ["renewable energy", "solar"],
    "CH-0025": ["eclipse", "moon"],
    "CH-0026": ["tết nguyên đán", "lunar"],
    "CH-0027": ["impressionism", "19th-century art movement"],
}
SOURCE_MISMATCH_ROWS = {"CH-0007", "CH-0016", *{f"CH-{n:04d}" for n in range(31, 51)}}
CURRENT_OR_HIGH_STAKES_ROWS = {
    "CH-0001", "CH-0003", "CH-0004", "CH-0009", "CH-0010", "CH-0012", "CH-0013",
    "CH-0014", "CH-0015", "CH-0019", "CH-0020", "CH-0021", "CH-0022", "CH-0023",
    "CH-0024", "CH-0028", "CH-0029", "CH-0030",
}
OFFICIAL_SOURCES = {
    "CH-0003": ("https://e-services.mps.gov.vn/bocongan/tintuc/chitiet?matin=42", "MPS public service portal", "official_government"),
    "CH-0013": ("https://vanban.chinhphu.vn/?pageid=27160&docid=198540", "Bộ Luật Lao động 45/2019/QH14", "official_government"),
    "CH-0023": ("https://vanban.chinhphu.vn/?pageid=27160&docid=212167&classid=1&orggroupid=2", "Nghị định 168/2024/NĐ-CP", "official_government"),
}


def digest(data: bytes | str) -> str:
    if isinstance(data, str):
        data = data.encode("utf-8")
    return "sha256:" + hashlib.sha256(data).hexdigest()


def normalize(value: str) -> str:
    return re.sub(r"\s+", " ", value.casefold()).strip()


def wiki_api_url(source_url: str) -> str | None:
    parsed = urlparse(source_url)
    if not parsed.hostname or not parsed.hostname.endswith("wikipedia.org"):
        return None
    language = parsed.hostname.split(".")[0]
    title = unquote(parsed.path.rsplit("/", 1)[-1])
    return f"https://{language}.wikipedia.org/api/rest_v1/page/summary/{quote(title, safe='()')}" if title else None


def first_sentence(text: str) -> str:
    match = re.match(r"(.+?[.!?。！？])(?:\s|$)", text.strip(), re.S)
    return (match.group(1) if match else text[:320]).strip()


def fetch_source(row: dict, cache_dir: Path) -> dict:
    source_url = str(row.get("gold_source_url") or "")
    api_url = wiki_api_url(source_url)
    result = {
        "source_fetch_status": "NOT_FETCHED_POLICY",
        "source_http_status": None,
        "source_fetch_timestamp": datetime.now(timezone.utc).isoformat(),
        "source_content_sha256": None,
        "source_cache_path": None,
        "source_api_url": api_url,
        "source_title_observed": row.get("gold_source_title", ""),
        "source_canonical_url_observed": source_url,
        "source_content_text": "",
        "source_fetch_error": None,
    }
    if not api_url:
        return result
    try:
        response = requests.get(api_url, headers={"User-Agent": "SCP-Phase3-Gold-Builder/2.0"}, timeout=(10, 30))
        result["source_http_status"] = response.status_code
        result["source_content_sha256"] = digest(response.content)
        cache_dir.mkdir(parents=True, exist_ok=True)
        cache_path = cache_dir / f"{row['question_id']}.source.json"
        cache_path.write_bytes(response.content)
        result["source_cache_path"] = str(cache_path)
        if response.status_code != 200:
            result["source_fetch_status"] = "FETCH_FAILED_HTTP"
            result["source_fetch_error"] = f"HTTP {response.status_code}"
            return result
        payload = response.json()
        text = str(payload.get("extract") or "").strip()
        if not text:
            result["source_fetch_status"] = "FETCHED_NO_TEXT"
            result["source_fetch_error"] = "empty summary extract"
            return result
        result["source_fetch_status"] = "FETCHED_TEXT"
        result["source_content_text"] = text
        result["source_title_observed"] = payload.get("title") or result["source_title_observed"]
        result["source_canonical_url_observed"] = (((payload.get("content_urls") or {}).get("desktop") or {}).get("page") or source_url)
    except (requests.RequestException, ValueError) as exc:
        result["source_fetch_status"] = "FETCH_FAILED_EXCEPTION"
        result["source_fetch_error"] = f"{type(exc).__name__}: {str(exc)[:240]}"
    return result


def classify(qid: str, source: dict) -> tuple[str, str, str, str]:
    if qid in SOURCE_MISMATCH_ROWS:
        return "SOURCE_MISMATCH_NEEDS_REBUILD", "MISMATCH_OBSERVED", "TEMPORAL_OR_STATIC_UNPROVEN", "Legacy source is unrelated to the question; rebuild required."
    if qid in CURRENT_OR_HIGH_STAKES_ROWS:
        return "HUMAN_REVIEW_REQUIRED_CURRENT_OR_HIGH_STAKES", "SOURCE_NOT_ACCEPTED_FOR_GOLD", "CURRENT_CLAIM_REQUIRES_FRESHNESS", "Use authoritative current source and human review; no automatic promotion."
    if qid in STATIC_WIKI_ROWS and source["source_fetch_status"] == "FETCHED_TEXT":
        return "CANDIDATE_SOURCE_FETCH_VERIFIED", "MACHINE_QUOTE_CONTAINED_ONLY", "STATIC_CORE_ONLY", "Fetched quote is contained in source text; answer alignment and promotion remain unreviewed."
    if qid in STATIC_WIKI_ROWS:
        return "SOURCE_FETCH_FAILED", "SOURCE_UNAVAILABLE", "STATIC_CORE_ONLY", "Static source could not be independently fetched."
    return "HUMAN_REVIEW_REQUIRED", "SOURCE_NOT_ACCEPTED_FOR_GOLD", "TEMPORAL_SCOPE_UNPROVEN", "No safe automatic promotion rule."


def build_row(row: dict, cache_dir: Path) -> dict:
    qid = row["question_id"]
    source = fetch_source(row, cache_dir) if qid in STATIC_WIKI_ROWS else {
        "source_fetch_status": "NOT_FETCHED_POLICY", "source_http_status": None,
        "source_fetch_timestamp": datetime.now(timezone.utc).isoformat(), "source_content_sha256": None,
        "source_cache_path": None, "source_api_url": None, "source_title_observed": row.get("gold_source_title", ""),
        "source_canonical_url_observed": row.get("gold_source_url", ""), "source_content_text": "", "source_fetch_error": None,
    }
    proposed_url, proposed_title, authority = OFFICIAL_SOURCES.get(qid, (row.get("gold_source_url", ""), row.get("gold_source_title", ""), "secondary_or_unknown"))
    status, alignment, temporal_scope, notes = classify(qid, source)
    text = source["source_content_text"]
    quote = first_sentence(text) if status == "CANDIDATE_SOURCE_FETCH_VERIFIED" else ""
    terms = STATIC_WIKI_ROWS.get(qid, [])
    matched = [term for term in terms if normalize(term) in normalize(text)]
    return {
        "schema_version": "phase3-gold-v2-candidate",
        "question_id": qid,
        "question": row.get("question", ""),
        "candidate_answer": row.get("gold_answer", ""),
        "legacy_source_url": row.get("gold_source_url", ""),
        "legacy_source_title": row.get("gold_source_title", ""),
        "proposed_canonical_source_url": proposed_url,
        "proposed_source_title": proposed_title,
        "source_authority": authority,
        "source_fetch_status": source["source_fetch_status"],
        "source_http_status": source["source_http_status"],
        "source_fetch_timestamp": source["source_fetch_timestamp"],
        "source_content_sha256": source["source_content_sha256"],
        "source_cache_path": source["source_cache_path"],
        "source_api_url": source["source_api_url"],
        "source_title_observed": source["source_title_observed"],
        "source_canonical_url_observed": source["source_canonical_url_observed"],
        "source_fetch_error": source["source_fetch_error"],
        "gold_chunk_id": f"phase3-v2-{qid}" if text else "",
        "gold_chunk_text": text,
        "gold_chunk_sha256": digest(text) if text else "",
        "gold_evidence_quote": quote,
        "quote_contained_in_fetched_source": bool(quote) and quote in text,
        "question_anchor_terms": terms,
        "matched_anchor_terms": matched,
        "alignment_status": alignment,
        "temporal_scope": temporal_scope,
        "review_status": status,
        "reviewer_id": "SCP-AUTO-FETCH-V3",
        "reviewer_provenance": "programmatic fetch and containment check only; not human review",
        "reviewed_at": datetime.now(timezone.utc).isoformat(),
        "review_method": "independent_source_fetch_only",
        "review_notes": notes,
        "human_review_required": True,
        "gold_promotion": "BLOCKED_PENDING_HUMAN_REVIEW",
        "eligible_for_ragas": False,
        "candidate_only": True,
    }


def write_outputs(rows: list[dict], output_dir: Path) -> dict:
    output_dir.mkdir(parents=True, exist_ok=True)
    jsonl = output_dir / "gold_anchor_50_v2_candidates.jsonl"
    csv_path = output_dir / "gold_anchor_50_v2_review_ledger.csv"
    xlsx = output_dir / "gold_anchor_50_v2_review_ledger.xlsx"
    report_path = output_dir / "gold_anchor_50_v2_validation.json"
    jsonl.write_text("\n".join(json.dumps(row, ensure_ascii=False) for row in rows) + "\n", encoding="utf-8")
    fields = list(rows[0])
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=fields)
        writer.writeheader()
        for row in rows:
            writer.writerow({key: json.dumps(value, ensure_ascii=False) if isinstance(value, (list, dict)) else value for key, value in row.items()})
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Gold Candidate v2"
    sheet.append(fields)
    for row in rows:
        sheet.append([json.dumps(row[field], ensure_ascii=False) if isinstance(row[field], (list, dict)) else row[field] for field in fields])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.alignment = Alignment(wrap_text=True, vertical="top")
    for row_idx in range(2, sheet.max_row + 1):
        fill = PatternFill("solid", fgColor="E2F0D9" if sheet.cell(row_idx, fields.index("review_status") + 1).value == "CANDIDATE_SOURCE_FETCH_VERIFIED" else "FCE4D6")
        for cell in sheet[row_idx]:
            cell.fill = fill
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    for index, field in enumerate(fields, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min({"question": 52, "candidate_answer": 55, "gold_chunk_text": 65, "gold_evidence_quote": 55, "review_notes": 55}.get(field, 24), 70)
    workbook.save(xlsx)
    report = {
        "schema_version": "phase3-gold-v2-validation",
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "rows": len(rows),
        "input_sha256": None,
        "output_hashes": {name: digest(path.read_bytes()) for name, path in (("jsonl", jsonl), ("csv", csv_path), ("xlsx", xlsx))},
        "status_counts": {},
        "quote_contained_count": sum(bool(row["quote_contained_in_fetched_source"]) for row in rows),
        "eligible_for_ragas_count": sum(bool(row["eligible_for_ragas"]) for row in rows),
        "human_review_required_count": sum(bool(row["human_review_required"]) for row in rows),
        "gate_status": "BLOCKED_PENDING_INDEPENDENT_REVIEW",
        "legacy_gold_preserved": True,
        "provenance_note": "Candidate answer was never used as evidence quote; human review identity was not fabricated.",
    }
    for row in rows:
        report["status_counts"][row["review_status"]] = report["status_counts"].get(row["review_status"], 0) + 1
    report_path.write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    return {"jsonl": str(jsonl), "csv": str(csv_path), "xlsx": str(xlsx), "report": str(report_path), "report_data": report}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input", type=Path, default=Path(__file__).resolve().parent / "gold_anchor_50_v1.jsonl")
    parser.add_argument("--output-dir", type=Path, default=Path(__file__).resolve().parent)
    parser.add_argument("--cache-dir", type=Path, default=Path(".artifacts") / "phase3_source_cache_v2")
    parser.add_argument("--delay-seconds", type=float, default=0.5)
    args = parser.parse_args()
    rows = [json.loads(line) for line in args.input.read_text(encoding="utf-8").splitlines() if line.strip()]
    if len(rows) != 50 or len({row.get("question_id") for row in rows}) != 50:
        raise SystemExit("input must contain 50 unique question_id rows")
    built = []
    for index, row in enumerate(rows):
        built.append(build_row(row, args.cache_dir))
        if index + 1 < len(rows):
            time.sleep(max(args.delay_seconds, 0))
    result = write_outputs(built, args.output_dir)
    result["report_data"]["input_sha256"] = digest(args.input.read_bytes())
    Path(result["report"]).write_text(json.dumps(result["report_data"], ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(result["report_data"], ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
